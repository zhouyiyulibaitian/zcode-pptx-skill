"""Deck -> PPTX with python-pptx.

Everything lands in PowerPoint as native objects: text boxes with real runs and
paragraph properties, preset-geometry autoshapes, tables, native editable charts
(so the data can still be re-edited in PowerPoint), speaker notes and pictures.
Nothing is flattened to an image, which is what makes the exported file
"完全可编辑".

Two build modes:
  * theme mode  - our own layouts, restyled with the theme (HTML preview matches
                  this exactly).
  * template mode - the deck is written on top of an existing .pptx template, so
                  the template's masters, layouts, theme fonts/colours and any
                  page furniture (校徽、页眉页脚、背景) are preserved and the slides
                  are ordinary slides of that file.
"""

from __future__ import annotations

import copy
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import model as M

# --------------------------------------------------------------------------- #
# low level helpers
# --------------------------------------------------------------------------- #

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}

SHAPE_MAP = {
    "rect": MSO_SHAPE.RECTANGLE,
    "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "line": MSO_SHAPE.RECTANGLE,          # drawn as a thin rect for full control
    "arrow": MSO_SHAPE.RECTANGLE,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "chevron": MSO_SHAPE.CHEVRON,
    "pentagon": MSO_SHAPE.REGULAR_PENTAGON,
    "star5": MSO_SHAPE.STAR_5_POINT,
    "parallelogram": MSO_SHAPE.PARALLELOGRAM,
    "plus": MSO_SHAPE.MATH_PLUS,
    "arc": MSO_SHAPE.BLOCK_ARC,
    "ring": MSO_SHAPE.DONUT,
    "cloud": MSO_SHAPE.CLOUD,
}

CHART_MAP = {
    "column": (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.COLUMN_STACKED),
    "bar": (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.BAR_STACKED),
    "line": (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE_MARKERS),
    "area": (XL_CHART_TYPE.AREA, XL_CHART_TYPE.AREA_STACKED),
    "pie": (XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE),
    "doughnut": (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.DOUGHNUT),
    "radar": (XL_CHART_TYPE.RADAR, XL_CHART_TYPE.RADAR),
}


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string((hexstr or "#000000").lstrip("#").upper())


def _set_typeface(rPr, tag: str, face: str) -> None:
    """Write one of a:latin / a:ea / a:cs, keeping the schema's child order.

    python-pptx only models a:latin (via font.name) and has no accessor for a:ea,
    so CJK text set with `font.name` alone would render in the theme's East Asian
    font instead of the one the preview used.
    """
    order = ['a:latin', 'a:ea', 'a:cs', 'a:sym', 'a:hlinkClick', 'a:hlinkMouseOver',
             'a:rtl', 'a:extLst']
    el = rPr.find(qn(tag))
    if el is None:
        el = rPr.makeelement(qn(tag), {'typeface': face})
        successors = order[order.index(tag) + 1:]
        rPr.insert_element_before(el, *successors)
    else:
        el.set('typeface', face)


def _set_font(run, font_spec: dict | None):
    """Apply latin + East Asian + complex-script typefaces to a run."""
    if not font_spec:
        return
    latin = font_spec.get("latin")
    ea = font_spec.get("ea") or latin
    rPr = run._r.get_or_add_rPr()
    if latin:
        _set_typeface(rPr, 'a:latin', latin)
    if ea:
        _set_typeface(rPr, 'a:ea', ea)
        _set_typeface(rPr, 'a:cs', ea)


def _solid_alpha(shape, alpha: float) -> None:
    """Add <a:alpha> to the shape's first solid fill (PowerPoint transparency)."""
    if alpha is None or float(alpha) >= 1:
        return
    spPr = shape._element.spPr
    for sf in spPr.findall(qn('a:solidFill')):
        srgb = sf.find(qn('a:srgbClr'))
        if srgb is not None and srgb.find(qn('a:alpha')) is None:
            srgb.append(srgb.makeelement(qn('a:alpha'),
                                        {'val': str(int(round(float(alpha) * 100000)))}))
            return


def _apply_fill(shape, fill: dict | None, fallback: str | None = None) -> None:
    if not fill:
        if fallback:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(fallback)
        else:
            shape.fill.background()
        return
    ftype = fill.get("type", "solid")
    if ftype == "gradient" and fill.get("from") and fill.get("to"):
        try:
            shape.fill.gradient()
            stops = shape.fill.gradient_stops
            stops[0].color.rgb = _rgb(fill["from"])
            stops[1].color.rgb = _rgb(fill["to"])
            if fill.get("angle") is not None:
                shape.fill.gradient_angle = float(fill.get("angle", 90))
            return
        except Exception:
            pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill.get("color") or fallback or "#FFFFFF")
    _solid_alpha(shape, fill.get("alpha"))


def _apply_line(shape, line: dict | None, dashed_ok: bool = True) -> None:
    if not line or not line.get("color"):
        shape.line.fill.background()
        return
    shape.line.color.rgb = _rgb(line["color"])
    shape.line.width = Pt(float(line.get("width", 1)))
    dash = line.get("dash")
    if dashed_ok and dash in ("dash", "dot", "lgDash"):
        shape.line.dash_style = {"dash": MSO_LINE_DASH_STYLE.DASH,
                                 "dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
                                 "lgDash": MSO_LINE_DASH_STYLE.LONG_DASH}[dash]
    if line.get("arrow"):
        ln = shape.line._get_or_add_ln()
        tail = ln.makeelement(qn('a:tailEnd'),
                              {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(tail)


def _set_adjust(shape, radius_in: float) -> None:
    """roundRect corner radius: adj is 1/100000 of the shorter side."""
    if not radius_in:
        return
    try:
        w = shape.width / M.EMU_PER_INCH
        h = shape.height / M.EMU_PER_INCH
        ratio = min(max(float(radius_in) / max(min(w, h), 0.01), 0.0), 0.5)
        prstGeom = shape._element.spPr.find(qn('a:prstGeom'))
        if prstGeom is None:
            return
        avLst = prstGeom.find(qn('a:avLst'))
        if avLst is None:
            avLst = prstGeom.makeelement(qn('a:avLst'), {})
            prstGeom.append(avLst)
        for old in avLst.findall(qn('a:gd')):
            avLst.remove(old)
        avLst.append(avLst.makeelement(qn('a:gd'),
                                       {'name': 'adj', 'fmla': f'val {int(ratio * 100000)}'}))
    except Exception:
        pass


def _set_bullet(p, bullet: bool, indent_in: float = 0.0) -> None:
    pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
    if indent_in:
        pPr.set('marL', str(M.inch_to_emu(indent_in)))
        pPr.set('indent', str(-M.inch_to_emu(indent_in)))
    if not bullet:
        if pPr.find(qn('a:buNone')) is None:
            pPr.append(pPr.makeelement(qn('a:buNone'), {}))
        return
    if not indent_in:
        pPr.set('marL', '228600')
        pPr.set('indent', '-228600')
    buFont = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
    buChar = pPr.makeelement(qn('a:buChar'), {'char': '•'})
    pPr.append(buFont)
    pPr.append(buChar)


def _set_every_paragraph_no_bullet(text_frame) -> None:
    for p in text_frame.paragraphs:
        _set_bullet(p, False)


# --------------------------------------------------------------------------- #
# text / shape / image / table / chart writers
# --------------------------------------------------------------------------- #

def _write_text_frame(tf, el: dict, theme: dict, scale: float = 1.0,
                      default_font: dict | None = None) -> None:
    body_font = default_font or (theme.get("fonts") or {}).get("body") or {}
    pad = M._padding(el)
    tf.word_wrap = bool(el.get("wrap", True))
    tf.margin_left = Inches(pad[0])
    tf.margin_right = Inches(pad[2])
    tf.margin_top = Inches(pad[1])
    tf.margin_bottom = Inches(pad[3])
    tf.vertical_anchor = ANCHOR.get(el.get("valign", "top"), MSO_ANCHOR.TOP)
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    paragraphs = el.get("paragraphs") or []
    for i, spec in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(spec.get("align", "left"), PP_ALIGN.LEFT)
        if spec.get("lineSpacing"):
            try:
                p.line_spacing = float(spec["lineSpacing"])
            except Exception:
                pass
        if spec.get("spaceBefore"):
            p.space_before = Pt(float(spec["spaceBefore"]) * scale)
        if spec.get("spaceAfter"):
            p.space_after = Pt(float(spec["spaceAfter"]) * scale)
        _set_bullet(p, bool(spec.get("bullet")), M._padding(el)[0] and 0.0 or 0.0)
        runs = spec.get("runs") or [{"text": ""}]
        for rspec in runs:
            r = p.add_run()
            r.text = str(rspec.get("text", ""))
            size = rspec.get("size") or spec.get("size")
            if size:
                r.font.size = Pt(float(size) * scale)
            if rspec.get("bold", spec.get("bold")):
                r.font.bold = True
            if rspec.get("italic", spec.get("italic")):
                r.font.italic = True
            if rspec.get("underline"):
                r.font.underline = True
            color = rspec.get("color") or spec.get("color")
            if color:
                r.font.color.rgb = _rgb(color)
            font = rspec.get("font") or spec.get("font") or body_font
            if font:
                if not isinstance(font, dict):
                    font = {"latin": font, "ea": font}
                _set_font(r, {"latin": font.get("latin") or body_font.get("latin"),
                              "ea": font.get("ea") or body_font.get("ea")})
            if rspec.get("spacing"):
                rPr = r._r.get_or_add_rPr()
                rPr.set('spc', str(int(float(rspec["spacing"]) * 100)))


def add_text(slide, el: dict, theme: dict) -> None:
    box = slide.shapes.add_textbox(Inches(el["x"]), Inches(el["y"]),
                                   Inches(el["w"]), Inches(el["h"]))
    box.name = f"Text: {el.get('role') or el.get('id')}"
    _write_text_frame(box.text_frame, el, theme)
    if el.get("rotation"):
        box.rotation = float(el["rotation"])
    if el.get("fill"):
        _apply_fill(box, el["fill"])
    if el.get("line"):
        _apply_line(box, el["line"])


def add_shape(slide, el: dict, theme: dict) -> None:
    shape_name = el.get("shape", "rect")
    if shape_name in ("line", "arrow"):
        # a hairline rectangle keeps stroke width control identical to the HTML preview
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(el["x"]), Inches(el["y"]),
                                     Inches(el["w"]), Inches(max(el["h"], 0.008)))
        box.fill.solid()
        _apply_fill(box, el.get("fill") or {"color": (el.get("line") or {}).get("color", "#000000")})
        box.line.fill.background()
        if shape_name == "arrow":
            ln = el.get("line") or {}
            box.fill.fore_color.rgb = _rgb(ln.get("color", "#000000"))
    else:
        box = slide.shapes.add_shape(SHAPE_MAP.get(shape_name, MSO_SHAPE.RECTANGLE),
                                     Inches(el["x"]), Inches(el["y"]),
                                     Inches(el["w"]), Inches(el["h"]))
        _apply_fill(box, el.get("fill"), theme["colors"].get("surface"))
        _apply_line(box, el.get("line"))
        if shape_name == "roundRect":
            _set_adjust(box, float(el.get("radius") or 0))
        if shape_name == "ring":
            _set_adjust(box, 0.0)
    box.name = f"Shape: {shape_name}"
    try:
        box.shadow.inherit = False
    except Exception:
        pass
    if el.get("rotation"):
        box.rotation = float(el["rotation"])
    if el.get("paragraphs"):
        _write_text_frame(box.text_frame, el, theme)
    else:
        box.text_frame.text = ""
        _set_every_paragraph_no_bullet(box.text_frame)


def add_image(slide, el: dict, deck_dir: str, theme: dict) -> bool:
    src = el.get("src") or ""
    path = src if os.path.isabs(src) else os.path.join(deck_dir, src)
    if not src or not os.path.exists(path):
        _add_placeholder(slide, el, theme)
        return False
    x, y, w, h = el["x"], el["y"], el["w"], el["h"]
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    pic.name = f"Image: {os.path.basename(src)}"
    # fit: 'cover' -> crop the overflow, 'contain' -> letterbox inside the box
    try:
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
        if iw and ih:
            box_ar = w / h if h else 1.0
            img_ar = iw / ih
            if el.get("fit", "cover") == "cover":
                if img_ar > box_ar:      # too wide -> crop left/right
                    keep = box_ar / img_ar
                    crop = (1 - keep) / 2
                    pic.crop_left = crop
                    pic.crop_right = crop
                else:                     # too tall -> crop top/bottom
                    keep = img_ar / box_ar
                    crop = (1 - keep) / 2
                    pic.crop_top = crop
                    pic.crop_bottom = crop
            else:
                if img_ar > box_ar:
                    new_h = w / img_ar
                    pic.top = Inches(y + (h - new_h) / 2)
                    pic.height = Inches(new_h)
                else:
                    new_w = h * img_ar
                    pic.left = Inches(x + (w - new_w) / 2)
                    pic.width = Inches(new_w)
    except Exception:
        pass
    radius = float(el.get("radius") or 0)
    if radius > 0:
        try:
            prstGeom = pic._element.spPr.find(qn('a:prstGeom'))
            if prstGeom is not None:
                prstGeom.set('prst', 'roundRect')
                avLst = prstGeom.find(qn('a:avLst'))
                if avLst is None:
                    avLst = prstGeom.makeelement(qn('a:avLst'), {})
                    prstGeom.append(avLst)
                ratio = min(radius / max(min(w, h), 0.01), 0.5)
                avLst.append(avLst.makeelement(qn('a:gd'), {
                    'name': 'adj', 'fmla': f'val {int(ratio * 100000)}'}))
        except Exception:
            pass
    if el.get("opacity") is not None and float(el["opacity"]) < 1:
        try:
            solidFill = pic._element.spPr.find(qn('a:solidFill'))
            if solidFill is not None:
                _solid_alpha(pic, float(el["opacity"]))
        except Exception:
            pass
    return True


def _add_placeholder(slide, el: dict, theme: dict) -> None:
    """A missing image must never break the build -- draw a labelled grey box."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(el["x"]), Inches(el["y"]),
                                 Inches(el["w"]), Inches(el["h"]))
    box.fill.solid()
    box.fill.fore_color.rgb = _rgb(theme["colors"].get("surface", "#EEEEEE"))
    box.line.color.rgb = _rgb(theme["colors"].get("line", "#CCCCCC"))
    box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    box.text_frame.text = el.get("placeholder") or "图片占位"
    for p in box.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        _set_bullet(p, False)
        for r in p.runs:
            r.font.size = Pt(12)
            r.font.color.rgb = _rgb(theme["colors"].get("muted", "#888888"))
    box.name = "Image placeholder"


def _cell_borders(cell, color: str, width_pt: float = 0.75) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        ln = tcPr.makeelement(qn(tag), {'w': str(int(width_pt * 12700)), 'cap': 'flat',
                                        'cmpd': 'sng', 'algn': 'ctr'})
        fill = ln.makeelement(qn('a:solidFill'), {})
        fill.append(fill.makeelement(qn('a:srgbClr'), {'val': (color or '#CCCCCC').lstrip('#').upper()}))
        ln.append(fill)
        tcPr.append(ln)


def add_table(slide, el: dict, theme: dict) -> None:
    cols = el.get("cols") or []
    rows = el.get("rows") or []
    ncol = len(cols) or (len(rows[0]) if rows else 1)
    nrow = len(rows) or 1
    gf = slide.shapes.add_table(nrow, ncol, Inches(el["x"]), Inches(el["y"]),
                                Inches(el.get("w", 6)), Inches(el.get("h", 2)))
    table = gf.table
    gf.name = "Table"
    tbl = gf._element.graphic.graphicData.tbl
    # PowerPoint reads the header/banding flags from a:tblPr, not a:tbl, so the
    # default "Medium Style 2" would otherwise tint cells we filled explicitly.
    for node in (tbl, tbl.find(qn('a:tblPr'))):
        if node is not None:
            node.set('firstRow', '0')
            node.set('bandRow', '0')
    total = float(el.get("w", 6))
    if cols:
        for i, cw in enumerate(cols[:ncol]):
            table.columns[i].width = Inches(float(cw))
    else:
        for i in range(ncol):
            table.columns[i].width = Inches(total / ncol)
    row_h = float(el.get("rowHeight", 0.45))
    for i in range(nrow):
        table.rows[i].height = Inches(row_h)
    header = bool(el.get("header", True))
    body_size = theme.get("geometry", {}).get("bodySize", 14) * 0.92
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            cell = table.cell(ri, ci)
            spec = row[ci] if ci < len(row) else {}
            if not isinstance(spec, dict):
                spec = {"text": str(spec)}
            is_head = header and ri == 0
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = ANCHOR.get(spec.get("valign", "middle"), MSO_ANCHOR.MIDDLE)
            cell.fill.solid()
            if spec.get("fill"):
                cell.fill.fore_color.rgb = _rgb(spec["fill"])
            elif is_head:
                cell.fill.fore_color.rgb = _rgb(theme["colors"].get("primary", "#333333"))
            else:
                cell.fill.fore_color.rgb = _rgb(theme["colors"].get("surface", "#FFFFFF") if
                                                ri % 2 else theme["colors"].get("bg", "#FFFFFF"))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = ALIGN.get(spec.get("align", "center" if is_head else "left"), PP_ALIGN.LEFT)
            _set_bullet(p, False)
            r = p.add_run()
            r.text = str(spec.get("text", ""))
            r.font.size = Pt(float(spec.get("size") or body_size))
            r.font.bold = bool(spec.get("bold", is_head))
            r.font.color.rgb = _rgb(spec.get("color") or
                                    (theme["colors"].get("onPrimary", "#FFFFFF") if is_head
                                     else theme["colors"].get("text", "#111111")))
            _set_font(r, theme.get("fonts", {}).get("body"))
            _cell_borders(cell, theme["colors"].get("line", "#CCCCCC"))


def add_chart(slide, el: dict, theme: dict) -> None:
    ctype = el.get("chart", "column")
    stacked = bool(el.get("stacked"))
    cats = [str(c) for c in (el.get("categories") or [])]
    series = el.get("series") or []
    if ctype == "scatter":
        data = XyChartData()
        for s in series:
            sd = data.add_series(str(s.get("name") or ""))
            for i, v in enumerate(s.get("values") or []):
                x = _float_or(cats[i] if i < len(cats) else i, i)
                sd.add_data_point(x, _float_or(v, 0.0))
        chart_type = XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS if el.get("lines") else XL_CHART_TYPE.XY_SCATTER
    else:
        data = CategoryChartData()
        data.categories = cats or [str(i + 1) for i in range(
            max((len(s.get("values") or []) for s in series), default=1))]
        for s in series:
            data.add_series(str(s.get("name") or ""),
                            tuple(_float_or(v, 0.0) if v is not None else None
                                  for v in (s.get("values") or [])))
        chart_type = CHART_MAP.get(ctype, CHART_MAP["column"])[1 if stacked else 0]

    gf = slide.shapes.add_chart(chart_type, Inches(el["x"]), Inches(el["y"]),
                                Inches(el["w"]), Inches(el["h"]), data)
    gf.name = f"Chart: {ctype}"
    chart = gf.chart
    chart.font.size = Pt(float(el.get("fontSize") or theme.get("geometry", {}).get("smallSize", 11)))
    chart.font.name = (theme.get("fonts", {}).get("body") or {}).get("latin", "Arial")
    if el.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = str(el["title"])
        for p in chart.chart_title.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = _rgb(theme["colors"].get("text", "#111111"))
    legend_on = el.get("legend")
    if legend_on is None:
        legend_on = len(series) >= 2
    chart.has_legend = bool(legend_on)
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False

    colors = el.get("colors") or _series_colors(theme, len(series))
    try:
        plot = chart.plots[0]
        plot.gap_width = int(el.get("gapWidth") or 60)
        if ctype in ("column", "bar") and stacked:
            plot.overlap = 100
        if el.get("dataLabels"):
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format = str(el.get("valueFormat") or "0.##")
            dl.number_format_is_linked = False
            dl.font.size = Pt(float(theme.get("geometry", {}).get("smallSize", 11)) * 0.9)
            dl.font.color.rgb = _rgb(theme["colors"].get("muted", "#555555"))
        for i, s in enumerate(plot.series):
            color = (series[i].get("color") if i < len(series) else None) or \
                    (colors[i] if i < len(colors) else None)
            if not color:
                continue
            if ctype in ("pie", "doughnut"):
                for j, pt in enumerate(s.points):
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = _rgb(
                        (series[i].get("colors") or [])[j] if j < len(series[i].get("colors") or [])
                        else colors[j % len(colors)])
                s.format.line.color.rgb = _rgb(theme["colors"].get("bg", "#FFFFFF"))
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = _rgb(color)
                if ctype in ("line", "radar"):
                    s.format.line.color.rgb = _rgb(color)
                    s.format.line.width = Pt(2.25)
                    try:
                        s.smooth = False
                    except Exception:
                        pass
                elif ctype == "area":
                    s.format.line.color.rgb = _rgb(color)
    except Exception:
        pass
    try:
        va = chart.value_axis
        va.has_major_gridlines = bool(el.get("grid", True))
        va.visible = bool(el.get("yAxis", True))
        va.tick_labels.font.size = Pt(float(theme.get("geometry", {}).get("smallSize", 11)) * 0.9)
        if el.get("yMin") is not None:
            va.minimum_scale = float(el["yMin"])
        if el.get("yMax") is not None:
            va.maximum_scale = float(el["yMax"])
        if el.get("valueFormat"):
            va.tick_labels.number_format = str(el["valueFormat"])
            va.tick_labels.number_format_is_linked = False
        ca = chart.category_axis
        ca.has_major_gridlines = False
        ca.tick_labels.font.size = Pt(float(theme.get("geometry", {}).get("smallSize", 11)) * 0.9)
        ca.major_tick_mark = XL_TICK_MARK.NONE
    except Exception:
        pass


def _float_or(v, default: float) -> float:
    try:
        if v is None or v == "":
            return default
        return float(v)
    except Exception:
        return default


def _series_colors(theme: dict, n: int) -> list[str]:
    base = [theme["colors"].get("primary", "#1F4E79"),
            theme["colors"].get("accent", "#C9A227"),
            theme["colors"].get("muted", "#5C6A80"),
            "#8FB0D9", "#D98F8F", "#7FC3A8", "#C2A3D9", "#D9C27F"]
    return [base[i % len(base)] for i in range(max(n, 1))]


# --------------------------------------------------------------------------- #
# slide / deck assembly
# --------------------------------------------------------------------------- #

def _add_transition(slide, kind: str) -> None:
    if not kind:
        return
    sld = slide._element
    for old in sld.findall(qn('p:transition')):
        sld.remove(old)
    tr = sld.makeelement(qn('p:transition'), {'spd': 'med'})
    if kind in ("fade", "cut"):
        tr.append(tr.makeelement(qn('p:fade'), {}))
    elif kind == "push":
        tr.append(tr.makeelement(qn('p:push'), {'dir': 'u'}))
    elif kind == "wipe":
        tr.append(tr.makeelement(qn('p:wipe'), {'dir': 'l'}))
    else:
        tr.append(tr.makeelement(qn('p:fade'), {}))
    sld.insert_element_before(tr, 'p:timing', 'p:extLst')


def _set_background(slide, bg: dict | None, theme: dict, deck_dir: str) -> None:
    if not bg:
        return
    t = bg.get("type", "solid")
    try:
        if t == "gradient" and bg.get("from") and bg.get("to"):
            slide.background.fill.gradient()
            stops = slide.background.fill.gradient_stops
            stops[0].color.rgb = _rgb(bg["from"])
            stops[1].color.rgb = _rgb(bg["to"])
            slide.background.fill.gradient_angle = float(bg.get("angle", 90))
        elif t == "image" and bg.get("src"):
            path = bg["src"] if os.path.isabs(bg["src"]) else os.path.join(deck_dir, bg["src"])
            if os.path.exists(path):
                W, H = slide.part.package.presentation_part.presentation.slide_width, \
                       slide.part.package.presentation_part.presentation.slide_height
                pic = slide.shapes.add_picture(path, 0, 0, W, H)
                slide.shapes._spTree.remove(pic._element)
                slide.shapes._spTree.insert(2, pic._element)
        else:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(bg.get("color") or theme["colors"]["bg"])
        return
    except Exception:
        pass
    try:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(bg.get("color") or theme["colors"]["bg"])
    except Exception:
        pass


def _remove_empty_placeholders(slide) -> None:
    """Template slides come with inherited placeholders; we place our own boxes."""
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)


def _clear_slides(prs) -> int:
    """Drop the template's own slides, keeping masters, layouts and theme.

    A .pptx used as a template almost always carries sample slides; writing our
    deck on top of them would silently produce a file with both.
    """
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    for sldId in ids:
        rId = sldId.get(qn('r:id'))
        try:
            prs.part.drop_rel(rId)
        except Exception:
            pass
        sldIdLst.remove(sldId)
    return len(ids)


def build(deck: dict, theme: dict, out_path: str, base_template: str | None = None,
          layout_choice: dict | None = None, transition: str | None = None,
          keep_template_slides: bool = False) -> dict:
    """Write the deck and return a small build report."""
    resolved = M.resolve_deck(deck, theme)
    W, H = M.slide_size(resolved)
    deck_dir = os.path.dirname(os.path.abspath(out_path))
    using_template = bool(base_template and os.path.exists(base_template or ""))
    if using_template:
        prs = Presentation(base_template)
        # keep the template's slide size -- a school template is often 4:3
        if not keep_template_slides:
            _clear_slides(prs)
        tpl_layouts = list(prs.slide_layouts)
    else:
        prs = Presentation()
        prs.slide_width = Inches(W)
        prs.slide_height = Inches(H)
        tpl_layouts = list(prs.slide_layouts)
    blank_layout = _pick_blank_layout(tpl_layouts) if using_template else prs.slide_layouts[6]

    report = {"slides": 0, "shapes": 0, "images": 0, "missing_images": [],
              "charts": 0, "tables": 0, "notes": 0, "size": None}

    for idx, slide_spec in enumerate(resolved.get("slides", [])):
        layout = blank_layout
        if using_template and layout_choice:
            key = slide_spec.get("layout") or ""
            if layout_choice.get(key):
                layout = _layout_by_name(tpl_layouts, layout_choice[key]) or blank_layout
        slide = prs.slides.add_slide(layout)
        if using_template:
            _remove_empty_placeholders(slide)
        _set_background(slide, slide_spec.get("background"), theme, deck_dir)
        els = sorted(slide_spec.get("elements") or [], key=lambda e: float(e.get("z", 0)))
        for el in els:
            etype = el.get("type")
            try:
                if etype == "text":
                    add_text(slide, el, theme)
                elif etype == "shape":
                    add_shape(slide, el, theme)
                elif etype == "image":
                    if not add_image(slide, el, deck_dir, theme):
                        report["missing_images"].append(el.get("src"))
                    else:
                        report["images"] += 1
                elif etype == "table":
                    add_table(slide, el, theme)
                    report["tables"] += 1
                elif etype == "chart":
                    add_chart(slide, el, theme)
                    report["charts"] += 1
                report["shapes"] += 1
            except Exception as exc:  # one bad element must not lose the whole deck
                slide_spec.setdefault("warnings", []).append(f"{el.get('id')}: {exc}")
                report.setdefault("errors", []).append(f"第{idx + 1}页 {el.get('id')}: {exc}")
        notes = (slide_spec.get("notes") or "").strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
            report["notes"] += 1
        _add_transition(slide, slide_spec.get("transition") or transition)
        report["slides"] += 1

    cp = prs.core_properties
    cp.title = resolved["deck"].get("title") or ""
    if resolved["deck"].get("meta", {}).get("author"):
        cp.author = resolved["deck"]["meta"]["author"]
    cp.comments = "由 pptx-studio 生成（HTML 可编辑工作流）"
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    prs.save(out_path)
    report["size"] = [W, H]
    report["template"] = base_template if using_template else None
    return report


def _pick_blank_layout(layouts):
    for lay in layouts:
        name = (lay.name or "").lower()
        if name in ("blank", "空白", "空"):
            return lay
    best, best_n = layouts[0], 99
    for lay in layouts:
        n = len(list(lay.placeholders))
        if n < best_n:
            best, best_n = lay, n
    return best


def _layout_by_name(layouts, name: str):
    for lay in layouts:
        if lay.name == name:
            return lay
    for lay in layouts:
        if name and name.lower() in (lay.name or "").lower():
            return lay
    return None
